"""
ASK MDR Workflow Data Processor
Integrates:
1. MM-ASK-1A-GEN01-ENG-MDR-0001_B1*.xlsx (5 tabs: WP01 TOPSIDE, WP01 JACKET, WP02 TOPSIDE, WP02 JACKET, PRO.ENGINEERING_MR TBE)
2. ASK Technical Document Log Report*.xlsx
Generates: data/ask/mdr_data.js
"""

import os
import sys
import glob
import json
import re
from pathlib import Path
from datetime import datetime, timedelta, date
import openpyxl
import pandas as pd

def calc_working_days(d1, d2):
    if d1 is None or d2 is None: return None
    if d2 < d1: return 0
    cur = d1 + timedelta(days=1)
    work_days = 0
    while cur <= d2:
        if cur.weekday() < 5: work_days += 1
        cur += timedelta(days=1)
    return work_days

MONTH_MAP = {
    'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'may': 5, 'jun': 6,
    'jul': 7, 'aug': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dec': 12
}

def extract_file_date_obj(filename):
    match = re.search(r'(\d{1,2})[\s\-_]+([A-Za-z]{3})[\s\-_]+(\d{2,4})', filename)
    if match:
        day = int(match.group(1))
        month_str = match.group(2).lower()
        year = int(match.group(3))
        if year < 100: year += 2000
        month = MONTH_MAP.get(month_str, 1)
        return date(year, month, day)
    return date.min

def parse_date_val(val):
    if val is None: return None
    if isinstance(val, (datetime, date)):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if s.upper() in ['N/A', '-', '', 'NONE', 'NULL']:
        return 'N/A'
    for fmt in ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d-%b-%Y']:
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except:
            pass
    return s

def process_ask_mdr(mdr_path=None, log_path=None, output_dir="data/ask"):
    if not mdr_path:
        files = glob.glob("archive_ask/*MDR*.xlsx") + glob.glob("*MDR*.xlsx")
        valid = [f for f in files if not os.path.basename(f).startswith("~$")]
        if not valid:
            print("Error: No ASK MDR file found.")
            return False
        valid.sort(key=lambda x: extract_file_date_obj(os.path.basename(x)), reverse=True)
        mdr_path = valid[0]

    if not log_path:
        files = glob.glob("archive_ask/*ASK Technical Document Log Report*.xlsx") + glob.glob("*ASK Technical Document Log Report*.xlsx")
        valid = [f for f in files if not os.path.basename(f).startswith("~$")]
        if not valid:
            print("Error: No ASK Technical Log found.")
            return False
        valid.sort(key=lambda x: extract_file_date_obj(os.path.basename(x)), reverse=True)
        log_path = valid[0]

    print(f"Loading ASK Technical Log: {log_path}...")
    wb_log = openpyxl.load_workbook(log_path, data_only=True)
    sheet_log = wb_log['Page 1']
    log_vals = sheet_log.values
    for _ in range(9): next(log_vals)
    headers_log = next(log_vals)
    df_log = pd.DataFrame(log_vals, columns=headers_log)
    df_log = df_log.dropna(subset=['Document No.']).copy()

    doc_log_map = {}
    for _, row in df_log.iterrows():
        d_no = str(row['Document No.'] or '').strip().replace('\u2010', '-')
        if not d_no: continue
        d_key = d_no.upper()
        if d_key not in doc_log_map: doc_log_map[d_key] = []

        in_str = str(row['Actual Submit In'] or '').strip()
        out_str = str(row['Actual Submit Out'] or '').strip()
        in_dt, out_dt = None, None
        dur = None
        if in_str:
            try: in_dt = datetime.strptime(in_str, '%d-%b-%Y %H:%M:%S')
            except: pass
        if out_str:
            try: out_dt = datetime.strptime(out_str, '%d-%b-%Y %H:%M:%S')
            except: pass
        if in_dt and out_dt:
            dur = calc_working_days(in_dt.date(), out_dt.date())

        doc_log_map[d_key].append({
            'rev': str(row['Rev.'] or '').strip(),
            'submittal_id': str(row['Submittal ID'] or '').strip(),
            'doc_category': str(row.get('Document Category') or '').strip(),
            'in_dt': in_dt.strftime('%Y-%m-%d') if in_dt else None,
            'out_dt': out_dt.strftime('%Y-%m-%d') if out_dt else None,
            'duration_days': dur,
            'status_code': str(row['Return Status Code'] or '').strip(),
            'status_name': str(row['Return Status Name'] or '').strip()
        })

    print(f"Indexed {len(doc_log_map)} documents from ASK log.")

    print(f"Loading ASK MDR: {mdr_path}...")
    wb_mdr = openpyxl.load_workbook(mdr_path, data_only=True)
    sheets = ['WP01 TOPSIDE', 'WP01 JACKET', 'WP02 TOPSIDE', 'WP02 JACKET', 'PRO.ENGINEERING_MR TBE']

    records = []
    seen_docs = set()

    for sname in sheets:
        if sname not in wb_mdr.sheetnames: continue
        sheet = wb_mdr[sname]
        header_row = 11 if sname == 'PRO.ENGINEERING_MR TBE' else 13

        for r in range(header_row + 1, sheet.max_row + 1):
            c_doc = sheet.cell(row=r, column=9)
            doc_no_val = c_doc.value
            if not doc_no_val: continue
            # Exclude strikethrough if present
            if bool(c_doc.font and c_doc.font.strike): continue
            
            doc_no = str(doc_no_val).strip().replace('\u2010', '-')
            if not doc_no.startswith('MM-'): continue
            
            d_upper = doc_no.upper()
            if d_upper in seen_docs: continue
            seen_docs.add(d_upper)

            title = str(sheet.cell(row=r, column=10).value or '').strip()
            disc = str(sheet.cell(row=r, column=4).value or 'GEN').strip()
            area = str(sheet.cell(row=r, column=3).value or 'GEN').strip()
            doc_type = str(sheet.cell(row=r, column=5).value or 'DOC').strip()

            if sname == 'PRO.ENGINEERING_MR TBE':
                ifr_plan = parse_date_val(sheet.cell(row=r, column=23).value) or parse_date_val(sheet.cell(row=r, column=22).value)
                ifa_plan = parse_date_val(sheet.cell(row=r, column=30).value) or parse_date_val(sheet.cell(row=r, column=29).value)
                afc_plan = parse_date_val(sheet.cell(row=r, column=36).value)
                ap_plan = parse_date_val(sheet.cell(row=r, column=42).value)
            else:
                ifr_plan = parse_date_val(sheet.cell(row=r, column=23).value) or parse_date_val(sheet.cell(row=r, column=22).value)
                ifa_plan = parse_date_val(sheet.cell(row=r, column=30).value) or parse_date_val(sheet.cell(row=r, column=29).value)
                afc_plan = parse_date_val(sheet.cell(row=r, column=37).value) or parse_date_val(sheet.cell(row=r, column=36).value)
                ap_plan = parse_date_val(sheet.cell(row=r, column=44).value) or parse_date_val(sheet.cell(row=r, column=43).value)

            logs = doc_log_map.get(d_upper, [])
            logs.sort(key=lambda x: x['in_dt'] or '9999-99-99')

            ifr_log = next((l for l in logs if l['rev'].upper().startswith('A')), None)
            ifa_log = next((l for l in logs if l['rev'].upper().startswith('B')), None)
            afc_log = next((l for l in logs if l['rev'].upper().startswith('C') or l['rev'].upper() in ['0', '00', '1']), None)

            if not ifr_log and len(logs) > 0: ifr_log = logs[0]
            if not ifa_log and ifa_plan not in ['N/A', None] and len(logs) > 1: ifa_log = logs[1]
            if not afc_log and len(logs) > (2 if ifa_log else 1): afc_log = logs[-1]

            # Incorp days
            ifa_incorp_days = None
            if ifa_log and ifa_log['in_dt'] and ifr_log and ifr_log['out_dt']:
                try:
                    dt_in = datetime.strptime(ifa_log['in_dt'], '%Y-%m-%d')
                    dt_out = datetime.strptime(ifr_log['out_dt'], '%Y-%m-%d')
                    ifa_incorp_days = calc_working_days(dt_out.date(), dt_in.date())
                except: pass

            afc_incorp_days = None
            prev_out = ifa_log['out_dt'] if (ifa_log and ifa_log['out_dt']) else (ifr_log['out_dt'] if ifr_log else None)
            if afc_log and afc_log['in_dt'] and prev_out:
                try:
                    dt_in = datetime.strptime(afc_log['in_dt'], '%Y-%m-%d')
                    dt_out = datetime.strptime(prev_out, '%Y-%m-%d')
                    afc_incorp_days = calc_working_days(dt_out.date(), dt_in.date())
                except: pass

            ap_incorp_days = None
            ap_out_dt = None
            if afc_log and afc_log['out_dt']:
                after_afc = [l for l in logs if l['in_dt'] and l['in_dt'] > afc_log['out_dt']]
                if after_afc:
                    try:
                        dt_in = datetime.strptime(after_afc[0]['in_dt'], '%Y-%m-%d')
                        dt_out = datetime.strptime(afc_log['out_dt'], '%Y-%m-%d')
                        diff = calc_working_days(dt_out.date(), dt_in.date())
                        if diff is not None and diff >= 0: ap_incorp_days = diff
                    except: pass
                    appr_log = next((l for l in after_afc if any(k in (l['status_code'] or '').upper() for k in ['APPR', 'CODE 1', 'CODE 2', 'AWC'])), None)
                    if appr_log: ap_out_dt = appr_log['out_dt']
                elif afc_log['status_code'] and any(k in afc_log['status_code'].upper() for k in ['APPR', 'CODE 1', 'CODE 2', 'AWC']):
                    ap_incorp_days = 0
                    ap_out_dt = afc_log['out_dt']

            cat = 'Part C - Engineering Dwg' if (any(k in doc_no for k in ['-DWG-', '-PID-', '-PFD-', '-UFD-', '-LAY-', '-SCH-', '-SLD-', '-ISO-']) or doc_type in ['DWG', 'PID', 'PFD', 'UFD', 'LAY', 'SCH', 'SLD', 'ISO', 'Drawing']) else 'Part B - Engineering Doc'
            is_approved = any('APPR' in (l['status_code'] or '').upper() for l in logs) or any('APPR' in (l['status_name'] or '').upper() for l in logs)

            work_pkg = sname if sname != 'PRO.ENGINEERING_MR TBE' else 'WP01 TOPSIDE'

            records.append({
                'doc_no': doc_no,
                'title': title,
                'contractor_set': work_pkg,
                'work_package': work_pkg,
                'discipline': disc,
                'area': area,
                'doc_type': doc_type,
                'doc_category': cat,

                'ifr': {
                    'plan': ifr_plan,
                    'rev': ifr_log['rev'] if ifr_log else None,
                    'submit_in': ifr_log['in_dt'] if ifr_log else None,
                    'submit_out': ifr_log['out_dt'] if ifr_log else None,
                    'review_days': ifr_log['duration_days'] if ifr_log else None,
                    'status_code': ifr_log['status_code'] if ifr_log else None,
                    'status_name': ifr_log['status_name'] if ifr_log else None,
                    'submittal_id': ifr_log['submittal_id'] if ifr_log else None
                },

                'ifa': {
                    'is_skipped': (ifa_plan in ['N/A', None] and cat == 'Part C - Engineering Dwg'),
                    'plan': ifa_plan,
                    'rev': ifa_log['rev'] if ifa_log else None,
                    'submit_in': ifa_log['in_dt'] if ifa_log else None,
                    'submit_out': ifa_log['out_dt'] if ifa_log else None,
                    'review_days': ifa_log['duration_days'] if ifa_log else None,
                    'incorp_days': ifa_incorp_days if (ifa_incorp_days is not None and ifa_incorp_days >= 0) else None,
                    'status_code': ifa_log['status_code'] if ifa_log else None,
                    'status_name': ifa_log['status_name'] if ifa_log else None,
                    'submittal_id': ifa_log['submittal_id'] if ifa_log else None
                },

                'afc': {
                    'plan': afc_plan,
                    'rev': afc_log['rev'] if afc_log else None,
                    'submit_in': afc_log['in_dt'] if afc_log else None,
                    'submit_out': afc_log['out_dt'] if afc_log else None,
                    'review_days': afc_log['duration_days'] if afc_log else None,
                    'incorp_days': afc_incorp_days if (afc_incorp_days is not None and afc_incorp_days >= 0) else None,
                    'status_code': afc_log['status_code'] if afc_log else None,
                    'status_name': afc_log['status_name'] if afc_log else None,
                    'submittal_id': afc_log['submittal_id'] if afc_log else None
                },

                'ap': {
                    'plan': ap_plan,
                    'is_approved': is_approved,
                    'code': 'APPR' if is_approved else '',
                    'incorp_days': ap_incorp_days,
                    'submit_out': ap_out_dt,
                    'latest_status': logs[-1]['status_code'] if logs else None
                },

                'total_revisions_logged': len(logs)
            })

    m = re.search(r'(\d{1,2})[-\s]([A-Za-z]{3})[-\s](\d{2,4})', os.path.basename(mdr_path))
    report_date = f"{m.group(1).zfill(2)}-{m.group(2).capitalize()}-{m.group(3)}" if m else "25-Sep-2026"

    os.makedirs(output_dir, exist_ok=True)
    out_file = Path(output_dir) / 'mdr_data.js'
    with open(out_file, 'w', encoding='utf-8') as f:
        f.write(f'window.MDR_REPORT_DATE = "{report_date}";\n')
        f.write(f'window.ASK_MDR_REPORT_DATE = "{report_date}";\n')
        f.write('window.MDR_DATA = ' + json.dumps(records, ensure_ascii=False) + ';')

    print(f"Successfully generated {out_file} with {len(records)} ASK deliverables (Report Date: {report_date}).")
    return True

if __name__ == '__main__':
    process_ask_mdr()
