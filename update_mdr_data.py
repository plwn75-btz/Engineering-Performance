"""
Z1F Multi-Contractor Deliverable (EDSR/MDR) Workflow Data Processor
Integrates the 4 sets of Contractor Engineering Deliverable Registers:
1. Jacket & Piles (WP1B): Z1F Jacket EDSR cut-off_*.xlsx
2. Topside Structure: Z1F Topsides EDSR cut-off_*.xlsx
3. Topside without Structure (00210): 00210 EDSR as of WE 42 *.xlsm
4. WP2 Pipelines & Subsea: Att.2. WP2 Engineering Progress Measurement_*.xlsx

Cross-referenced against:
- Z1F Technical Document Log Report - 16 Oct 26.xlsx

Generates: mdr_data.js for index_mdr.html
"""

import os
import re
import sys
import glob
import json
from pathlib import Path
from datetime import datetime, timedelta, date
import openpyxl
import pandas as pd

def calc_working_days(d1, d2):
    if d1 is None or d2 is None: return None
    if d2 < d1: return 0
    cur = d1 + timedelta(days=1)
    work_days = 0
    while cur <= d2:
        if cur.weekday() < 5:  # Monday to Friday (0-4)
            work_days += 1
        cur += timedelta(days=1)
    return work_days

MONTH_MAP = {
    'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'may': 5, 'jun': 6,
    'jul': 7, 'aug': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dec': 12
}

def extract_file_date_obj(filename):
    match = re.search(r'(\d{1,2})[\s\-_]+([A-Za-z]{3})[\s\-_]+(\d{2,4})', filename)
    if match:
        day = int(match.group(1))
        month_str = match.group(2).lower()
        year = int(match.group(3))
        if year < 100: year += 2000
        month = MONTH_MAP.get(month_str, 1)
        return date(year, month, day)
    return date.min

def parse_date_val(val):
    if val is None: return None
    if isinstance(val, (datetime, date)):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if s.upper() in ['N/A', '-', '', 'NONE', 'NULL']:
        return 'N/A'
    for fmt in ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d-%b-%Y']:
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except:
            pass
    return s

def process_mdr(tech_log_path=None):
    # Locate Technical Log
    if not tech_log_path:
        files = glob.glob("Z1F Technical Document Log Report*.xlsx") + glob.glob("16102026/Z1F Technical Document Log Report*.xlsx")
        valid = [f for f in files if not os.path.basename(f).startswith("~$")]
        if not valid:
            print("Error: No Z1F Technical Document Log Report found.")
            return False
        valid.sort(key=lambda x: extract_file_date_obj(os.path.basename(x)), reverse=True)
        tech_log_path = valid[0]

    print(f"Loading Technical Log for Transmittal Cross-Reference: {tech_log_path}...")
    wb_log = openpyxl.load_workbook(tech_log_path, data_only=True)
    sheet_log = wb_log['Page 1']
    log_vals = sheet_log.values
    for _ in range(9): next(log_vals)
    headers_log = next(log_vals)
    df_log = pd.DataFrame(log_vals, columns=headers_log)
    df_log = df_log.dropna(subset=['Document No.']).copy()

    doc_log_map = {}
    for _, row in df_log.iterrows():
        d_no = str(row['Document No.'] or '').strip().replace('\u2010', '-')
        if not d_no: continue
        d_key = d_no.upper()
        if d_key not in doc_log_map: doc_log_map[d_key] = []

        in_str = str(row['Actual Submit In'] or '').strip()
        out_str = str(row['Actual Submit Out'] or '').strip()
        in_dt, out_dt = None, None
        dur = None
        if in_str:
            try: in_dt = datetime.strptime(in_str, '%d-%b-%Y %H:%M:%S')
            except: pass
        if out_str:
            try: out_dt = datetime.strptime(out_str, '%d-%b-%Y %H:%M:%S')
            except: pass
        if in_dt and out_dt:
            dur = calc_working_days(in_dt.date(), out_dt.date())

        doc_log_map[d_key].append({
            'rev': str(row['Rev.'] or '').strip(),
            'submittal_id': str(row['Submittal ID'] or '').strip(),
            'doc_category': str(row.get('Document Category') or '').strip(),
            'in_dt': in_dt.strftime('%Y-%m-%d') if in_dt else None,
            'out_dt': out_dt.strftime('%Y-%m-%d') if out_dt else None,
            'duration_days': dur,
            'status_code': str(row['Return Status Code'] or '').strip(),
            'status_name': str(row['Return Status Name'] or '').strip()
        })

    print(f"Indexed {len(doc_log_map)} unique documents from Technical Transmittal Log.")

    # 1. Parse Jacket EDSR
    jacket_files = glob.glob("*Jacket*EDSR*.xlsx") + glob.glob("16102026/*Jacket*EDSR*.xlsx")
    j_records = []
    if jacket_files:
        print(f"Loading Jacket EDSR: {jacket_files[0]}...")
        wb_j = openpyxl.load_workbook(jacket_files[0], data_only=True)
        if 'EDSR' in wb_j.sheetnames:
            for r in wb_j['EDSR'].iter_rows(min_row=4):
                c_doc = r[2]
                if bool(c_doc.font and c_doc.font.strike): continue
                doc_no = str(c_doc.value or '').strip().replace('\u2010', '-')
                if not doc_no.startswith('MM-ZTK'): continue
                j_records.append({
                    'contractor_set': 'Jacket & Piles',
                    'work_package': 'WP1B Jacket & Piles',
                    'doc_no': doc_no,
                    'facility': str(r[0].value or '').strip(),
                    'type': str(r[1].value or '').strip(),
                    'title': str(r[3].value or '').strip(),
                    'ifr_plan': parse_date_val(r[9].value),
                    'ifa_plan': parse_date_val(r[12].value),
                    'afc_plan': parse_date_val(r[15].value),
                    'ap_plan': parse_date_val(r[18].value),
                    'status': str(r[21].value or '').strip()
                })

    # 2. Parse Topsides EDSR
    topsides_files = glob.glob("*Topsides*EDSR*.xlsx") + glob.glob("16102026/*Topsides*EDSR*.xlsx")
    t_records = []
    if topsides_files:
        print(f"Loading Topsides EDSR: {topsides_files[0]}...")
        wb_t = openpyxl.load_workbook(topsides_files[0], data_only=True)
        if 'EDSR' in wb_t.sheetnames:
            for r in wb_t['EDSR'].iter_rows(min_row=4):
                c_doc = r[3]
                if bool(c_doc.font and c_doc.font.strike): continue
                doc_no = str(c_doc.value or '').strip().replace('\u2010', '-')
                if not doc_no.startswith('MM-ZTK'): continue
                t_records.append({
                    'contractor_set': 'Topside Structure',
                    'work_package': 'Topside Structure',
                    'doc_no': doc_no,
                    'facility': str(r[0].value or '').strip(),
                    'type': str(r[2].value or '').strip(),
                    'title': str(r[4].value or '').strip(),
                    'ifr_plan': parse_date_val(r[10].value),
                    'ifa_plan': parse_date_val(r[13].value),
                    'afc_plan': parse_date_val(r[16].value),
                    'ap_plan': parse_date_val(r[19].value),
                    'status': str(r[22].value or '').strip()
                })

    # 3. Parse Topside without Structure (00210 EDSR)
    worley_files = glob.glob("*00210*EDSR*.xlsm") + glob.glob("16102026/*00210*EDSR*.xlsm")
    w_records = []
    if worley_files:
        print(f"Loading Topside without Structure (00210) EDSR: {worley_files[0]}...")
        wb_w = openpyxl.load_workbook(worley_files[0], read_only=True, data_only=True)
        # Parse Gates for milestones
        gates_map = {}
        if 'Gates' in wb_w.sheetnames:
            for r in wb_w['Gates'].iter_rows(min_row=2, values_only=True):
                did = str(r[0] or '').strip().upper()
                st = str(r[2] or '').strip().upper()
                if not did.startswith('MM-ZTK'): continue
                if did not in gates_map: gates_map[did] = {}
                gates_map[did][st] = {
                    'plan': parse_date_val(r[10]),
                    'actual': parse_date_val(r[12])
                }
        if 'DEL' in wb_w.sheetnames:
            for r in wb_w['DEL'].iter_rows(min_row=2, values_only=True):
                doc_no = str(r[0] or '').strip().replace('\u2010', '-')
                if not doc_no.startswith('MM-ZTK'): continue
                g = gates_map.get(doc_no.upper(), {})
                w_records.append({
                    'contractor_set': 'Topside without Structure',
                    'work_package': 'Topside without Structure',
                    'doc_no': doc_no,
                    'facility': str(r[73] or 'GEN').strip(),
                    'type': str(r[14] or '').strip(),
                    'title': str(r[1] or '').strip(),
                    'ifr_plan': g.get('IFR', {}).get('plan'),
                    'ifa_plan': g.get('IFA', {}).get('plan'),
                    'afc_plan': g.get('AFC', {}).get('plan'),
                    'ap_plan': g.get('AP', {}).get('plan'),
                    'status': str(r[11] or '').strip()
                })

    # 4. Parse WP2 Pipeline Progress
    wp2_files = glob.glob("*WP2*Progress*.xlsx") + glob.glob("16102026/*WP2*Progress*.xlsx")
    p_records = []
    if wp2_files:
        print(f"Loading WP2 Pipeline Progress: {wp2_files[0]}...")
        wb_p = openpyxl.load_workbook(wp2_files[0], read_only=True, data_only=True)
        if 'MDR-Detailed Design-A10' in wb_p.sheetnames:
            for r in wb_p['MDR-Detailed Design-A10'].iter_rows(min_row=11, values_only=True):
                doc_no = str(r[8] or '').strip().replace('\u2010', '-')
                if not doc_no.startswith('MM-ZTK'): continue
                p_records.append({
                    'contractor_set': 'WP2 Pipeline',
                    'work_package': 'WP2 Pipelines & Subsea',
                    'doc_no': doc_no,
                    'facility': str(r[2] or 'IPX').strip(),
                    'type': str(r[4] or '').strip(),
                    'title': str(r[9] or '').strip(),
                    'ifr_plan': parse_date_val(r[20]),
                    'ifa_plan': parse_date_val(r[29]),
                    'afc_plan': parse_date_val(r[36]),
                    'ap_plan': parse_date_val(r[41]),
                    'status': 'Active'
                })

    raw_deliverables = j_records + t_records + w_records + p_records
    print(f"Aggregated raw deliverables: Jacket={len(j_records)}, Topside Structure={len(t_records)}, Topside without Structure={len(w_records)}, WP2={len(p_records)} -> Total {len(raw_deliverables)}")

    # Deduplicate deliverables by doc_no (case-insensitive)
    seen_docs = {}
    for item in raw_deliverables:
        k = item['doc_no'].upper()
        if k not in seen_docs:
            seen_docs[k] = item
        else:
            # Merge missing plans if present
            for fld in ['ifr_plan', 'ifa_plan', 'afc_plan', 'ap_plan']:
                if not seen_docs[k].get(fld) and item.get(fld):
                    seen_docs[k][fld] = item[fld]

    print(f"Total Unique Deliverables across 4 Contractors: {len(seen_docs)}")

    records = []
    for doc_no_u, item in seen_docs.items():
        doc_no = item['doc_no']
        title = item['title']
        contractor_set = item['contractor_set']
        work_package = item['work_package']
        
        # Derive discipline, area, doc_type from document number tokens
        tokens = doc_no.split('-')
        area = tokens[3] if len(tokens) > 3 else item.get('facility', 'GEN')
        disc = tokens[4] if len(tokens) > 4 else 'GEN'
        doc_type = tokens[5] if len(tokens) > 5 else item.get('type', 'DOC')

        logs = doc_log_map.get(doc_no_u, [])
        logs.sort(key=lambda x: x['in_dt'] or '9999-99-99')

        ifr_plan = item.get('ifr_plan')
        ifa_plan = item.get('ifa_plan')
        afc_plan = item.get('afc_plan')
        ap_plan = item.get('ap_plan')

        # Identify log entries corresponding to revision stages
        ifr_log = next((l for l in logs if l['rev'].upper().startswith('A')), None)
        ifa_log = next((l for l in logs if l['rev'].upper().startswith('B')), None)
        afc_log = next((l for l in logs if l['rev'].upper().startswith('C') or l['rev'].upper() in ['0', '00', '1', '01']), None)

        if not ifr_log and len(logs) > 0: ifr_log = logs[0]
        if not ifa_log and ifa_plan not in ['N/A', None] and len(logs) > 1: ifa_log = logs[1]
        if not afc_log and len(logs) > (2 if (ifa_log or ifa_plan != 'N/A') else 1):
            afc_log = logs[-1]

        # Contractor incorporation turnaround (Mon-Fri Working Days)
        ifa_incorp_days = None
        if ifa_log and ifa_log['in_dt'] and ifr_log and ifr_log['out_dt']:
            try:
                dt_in = datetime.strptime(ifa_log['in_dt'], '%Y-%m-%d')
                dt_out = datetime.strptime(ifr_log['out_dt'], '%Y-%m-%d')
                ifa_incorp_days = calc_working_days(dt_out.date(), dt_in.date())
            except: pass

        afc_incorp_days = None
        prev_out = ifa_log['out_dt'] if (ifa_log and ifa_log['out_dt']) else (ifr_log['out_dt'] if ifr_log else None)
        if afc_log and afc_log['in_dt'] and prev_out:
            try:
                dt_in = datetime.strptime(afc_log['in_dt'], '%Y-%m-%d')
                dt_out = datetime.strptime(prev_out, '%Y-%m-%d')
                afc_incorp_days = calc_working_days(dt_out.date(), dt_in.date())
            except: pass

        # AFC -> AP close-out turnaround
        ap_incorp_days = None
        ap_out_dt = None
        if afc_log and afc_log['out_dt']:
            after_afc = [l for l in logs if l['in_dt'] and l['in_dt'] > afc_log['out_dt']]
            if after_afc:
                try:
                    dt_in = datetime.strptime(after_afc[0]['in_dt'], '%Y-%m-%d')
                    dt_out = datetime.strptime(afc_log['out_dt'], '%Y-%m-%d')
                    diff = calc_working_days(dt_out.date(), dt_in.date())
                    if diff is not None and diff >= 0:
                        ap_incorp_days = diff
                except: pass
                appr_log = next((l for l in after_afc if any(k in (l['status_code'] or '').upper() for k in ['APPR', 'CODE 1', 'CODE 2', 'AWC'])), None)
                if appr_log: ap_out_dt = appr_log['out_dt']
            elif afc_log['status_code'] and any(k in afc_log['status_code'].upper() for k in ['APPR', 'CODE 1', 'CODE 2', 'AWC']):
                ap_incorp_days = 0
                ap_out_dt = afc_log['out_dt']

        # Document Category (Engineering Doc vs Engineering Dwg)
        cat = 'Part C - Engineering Dwg' if (any(k in doc_no for k in ['-DWG-', '-PID-', '-PFD-', '-UFD-', '-LAY-', '-SCH-', '-SLD-', '-ISO-']) or doc_type in ['DWG', 'PID', 'PFD', 'UFD', 'LAY', 'SCH', 'SLD', 'ISO', 'Drawing']) else 'Part B - Engineering Doc'

        is_approved = any('APPR' in (l['status_code'] or '').upper() for l in logs) or any('APPR' in (l['status_name'] or '').upper() for l in logs)

        records.append({
            'doc_no': doc_no,
            'title': title,
            'contractor_set': contractor_set,
            'work_package': work_package,
            'discipline': disc,
            'area': area,
            'doc_type': doc_type,
            'doc_category': cat,

            'ifr': {
                'plan': ifr_plan,
                'rev': ifr_log['rev'] if ifr_log else None,
                'submit_in': ifr_log['in_dt'] if ifr_log else None,
                'submit_out': ifr_log['out_dt'] if ifr_log else None,
                'review_days': ifr_log['duration_days'] if ifr_log else None,
                'status_code': ifr_log['status_code'] if ifr_log else None,
                'status_name': ifr_log['status_name'] if ifr_log else None,
                'submittal_id': ifr_log['submittal_id'] if ifr_log else None
            },

            'ifa': {
                'is_skipped': (ifa_plan in ['N/A', None] and cat == 'Part C - Engineering Dwg'),
                'plan': ifa_plan,
                'rev': ifa_log['rev'] if ifa_log else None,
                'submit_in': ifa_log['in_dt'] if ifa_log else None,
                'submit_out': ifa_log['out_dt'] if ifa_log else None,
                'review_days': ifa_log['duration_days'] if ifa_log else None,
                'incorp_days': ifa_incorp_days if (ifa_incorp_days is not None and ifa_incorp_days >= 0) else None,
                'status_code': ifa_log['status_code'] if ifa_log else None,
                'status_name': ifa_log['status_name'] if ifa_log else None,
                'submittal_id': ifa_log['submittal_id'] if ifa_log else None
            },

            'afc': {
                'plan': afc_plan,
                'rev': afc_log['rev'] if afc_log else None,
                'submit_in': afc_log['in_dt'] if afc_log else None,
                'submit_out': afc_log['out_dt'] if afc_log else None,
                'review_days': afc_log['duration_days'] if afc_log else None,
                'incorp_days': afc_incorp_days if (afc_incorp_days is not None and afc_incorp_days >= 0) else None,
                'status_code': afc_log['status_code'] if afc_log else None,
                'status_name': afc_log['status_name'] if afc_log else None,
                'submittal_id': afc_log['submittal_id'] if afc_log else None
            },

            'ap': {
                'plan': ap_plan,
                'is_approved': is_approved,
                'code': 'APPR' if is_approved else '',
                'incorp_days': ap_incorp_days,
                'submit_out': ap_out_dt,
                'latest_status': logs[-1]['status_code'] if logs else None
            },

            'total_revisions_logged': len(logs)
        })

    report_date = "16-Oct-2026"
    out_file = Path('mdr_data.js')
    content = (
        f'window.MDR_REPORT_DATE = "{report_date}";\n'
        f'window.Z1F_MDR_REPORT_DATE = "{report_date}";\n'
        'window.MDR_DATA = ' + json.dumps(records, ensure_ascii=False) + ';'
    )
    with open(out_file, 'w', encoding='utf-8') as f:
        f.write(content)

    z1f_dir = Path("data/z1f")
    z1f_dir.mkdir(parents=True, exist_ok=True)
    with open(z1f_dir / "mdr_data.js", 'w', encoding='utf-8') as f:
        f.write(content)

    print(f"Successfully generated {out_file} and {z1f_dir / 'mdr_data.js'} with {len(records)} Z1F deliverables (Report Date: {report_date}).")
    
    # Contractor Breakdown
    print("\n--- Contractor Deliverables Breakdown in mdr_data.js ---")
    c_counts = pd.Series([r['contractor_set'] for r in records]).value_counts()
    print(c_counts)

    return True

if __name__ == '__main__':
    process_mdr()
